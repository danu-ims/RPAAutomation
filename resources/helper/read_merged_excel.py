from openpyxl import load_workbook
from robot.api.deco import keyword  # Import decorator keyword

def read_merged_excel(file_path, sheet_name, start_row=1):
    wb = load_workbook(file_path)
    sheet = wb[sheet_name]

    merged_cells_map = {}

    # Ambil semua merged cells
    for merged_range in sheet.merged_cells.ranges:
        top_left_cell = merged_range.min_row, merged_range.min_col
        merged_value = sheet.cell(*top_left_cell).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_cells_map[(row, col)] = merged_value

    data = []
    for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        new_row = []
        for col_idx, value in enumerate(row, start=1):
            if value is None and (row_idx, col_idx) in merged_cells_map:
                value = merged_cells_map[(row_idx, col_idx)]  # Isi dengan nilai merged cell
            new_row.append(value)
        data.append(new_row)

    return data[start_row - 1:]  # Return data mulai dari start_row
